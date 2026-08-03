#!/usr/bin/env python3
"""ADDITIVE writeback of automated-test results into a COPY of the unit-test workbook.

    original (never modified) : docs/OneDrive_1_8-1-2026/Report5_Unit Test.xlsx
    output                    : docs/audit/Report5_Unit_Test_v2.xlsx

This is deliberately NOT a rewrite. `Report5_Unit Test.xlsx` records MANUAL test execution;
the generated `*.uc.spec.ts` suite is AUTOMATED regression testing. They are different
artifacts and must not overwrite each other.

SURGICAL AT THE ZIP LEVEL
-------------------------
An earlier version loaded the workbook with openpyxl and re-saved it. openpyxl's writer only
round-trips what its object model understands, so that silently DROPPED real content:
xl/media/image1.png (an embedded image), xl/drawings/*, xl/webextensions/*,
xl/printerSettings/*, and a worksheet .rels part. Cell-level checks cannot detect that -- the
cells were all correct while the file was quietly damaged.

So this version never round-trips. It copies the original zip entry-by-entry and rewrites ONLY
the specific XML parts that must change, leaving every other entry byte-identical. openpyxl is
still used, but strictly READ-ONLY against the original, to compute *what* to change.

WHAT IT ADDS (per function sheet, as two new rows -- UTCIDs are columns in these decision
tables, so a per-UTCID "column" is structurally a row):
    "Automated test"  -- the generated test name, or BLANK if none exists yet
    "Spec alignment"  -- MATCHES / DIVERGES, or BLANK

WHAT IT CORRECTS (the only two things, both arithmetic, neither a manual observation):
  1. Statistics per-function tallies. 33 rows over-count bucket A by exactly +1 each, making
     the =SUM() totals read 431 instead of 398. Verified: in all 87 function sheets the UTCID
     column count == the manual Passed/Failed mark count == the parsed matrix count. So the
     function sheets and the manual log are already correct at 398; only the Statistics tally
     is wrong. Correcting it reconciles the tally TO the manual results -- it overwrites no
     manual judgement. Passed/Failed/Untested are re-derived from each sheet's own marks.
     The =SUM() formulas in row 99 are genuine and were always faithful; they are PRESERVED.
     Their cached <v> values are refreshed so non-Excel readers see the corrected totals too,
     and workbook.xml gets fullCalcOnLoad="1" so Excel recomputes on open regardless.
  2. The 11 MISMAPPED rows on the Functions sheet get the real Class.method(), taken from
     docs/audit/qa-recon-symbols.json (actualClass + the method parsed out of `signature` --
     that JSON has no `actualMethod` key, and claimedClass/claimedMethod are the WRONG
     spreadsheet claim).

WHAT IT NEVER TOUCHES:
    every per-UTCID manual result cell (the Passed/Failed and Executed Date rows in each
    function sheet), every condition/expected cell, every embedded image/drawing/setting,
    and the original file.

A UTCID with no automated test gets BLANK cells -- never "Failed", never "Untested".

IDEMPOTENT BY CONSTRUCTION: every run starts from the pristine original, so the added rows can
never duplicate and corrections can never compound. Safe to re-run at any coverage level;
running at 42 cases and later at 300 both produce a valid v2 workbook.

New string content is written as inline strings (t="inlineStr"), so xl/sharedStrings.xml stays
byte-identical rather than needing its count/uniqueCount rewritten.

Usage:  python3 scripts/qa/writeback-unittest-xlsx.py
"""

from __future__ import annotations

import json
import re
import shutil
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

import openpyxl
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[2]
SRC = REPO / "docs" / "OneDrive_1_8-1-2026" / "Report5_Unit Test.xlsx"
DST = REPO / "docs" / "audit" / "Report5_Unit_Test_v2.xlsx"
MATRIX = REPO / "docs" / "audit" / "uc-matrix.json"
RESULTS = REPO / "docs" / "audit" / "uc-results.json"
SYMBOLS = REPO / "docs" / "audit" / "qa-recon-symbols.json"
CHANGELOG_RAW = REPO / "docs" / "audit" / "xlsx-changelog-raw.txt"

META_SHEETS = {"Cover", "Functions", "Statistics"}
UTCID_HEADER_ROW = 9           # verified identical across all 87 function sheets
LABEL_COL = 2                  # column B holds "Passed/Failed", "Executed Date", "Defect ID"
ROW_AUTOMATED = "Automated test"
ROW_ALIGNMENT = "Spec alignment"

STATS_FIRST, STATS_LAST, STATS_SUM = 12, 98, 99
C_PASSED, C_FAILED, C_UNTESTED, C_N, C_A, C_B, C_TOTAL = 3, 4, 5, 6, 7, 8, 9
FUNC_FIRST, FUNC_LAST = 11, 97
F_CLASS, F_METHOD, F_CODE, F_SHEET = 3, 4, 5, 6

METHOD_RE = re.compile(r"(?:async\s+)?([A-Za-z_]\w*)\s*\(")


# --------------------------------------------------------------------------------------
# XML cell surgery. Operates on raw worksheet XML so that untouched bytes stay untouched.
# --------------------------------------------------------------------------------------

def _cell_re(ref: str) -> re.Pattern:
    return re.compile(r'<c r="' + re.escape(ref) + r'"((?:\s[^>]*?)?)(?:/>|>(.*?)</c>)', re.S)


def _strip_type(attrs: str) -> str:
    return re.sub(r'\s+t="[^"]*"', "", attrs)


def set_numeric(xml: str, ref: str, value) -> str:
    """Replace a cell's value, preserving its style and any <f> formula."""
    pat = _cell_re(ref)
    m = pat.search(xml)
    if not m:
        raise KeyError(f"cell {ref} not present in sheet XML")
    attrs, inner = _strip_type(m.group(1)), m.group(2) or ""
    fml = re.search(r"<f[^>]*>.*?</f>|<f[^>]*/>", inner, re.S)
    new_inner = (fml.group(0) if fml else "") + f"<v>{value}</v>"
    return xml[: m.start()] + f"<c r=\"{ref}\"{attrs}>{new_inner}</c>" + xml[m.end():]


def set_inline_string(xml: str, ref: str, text: str) -> str:
    """Set a cell to an inline string, preserving its style. Leaves sharedStrings.xml alone."""
    pat = _cell_re(ref)
    m = pat.search(xml)
    if not m:
        raise KeyError(f"cell {ref} not present in sheet XML")
    attrs = _strip_type(m.group(1))
    body = f'<is><t xml:space="preserve">{xml_escape(text)}</t></is>'
    return xml[: m.start()] + f'<c r="{ref}"{attrs} t="inlineStr">{body}</c>' + xml[m.end():]


def append_rows(xml: str, rows: list[tuple[int, dict[int, str], str | None]]) -> str:
    """Append <row> elements just before </sheetData>.

    rows: (row_number, {col_index: text}, style_index_or_None). Cells with empty text are
    omitted entirely -- an absent cell IS a blank cell in xlsx, which is exactly what a UTCID
    with no automated test should get.
    """
    if not rows:
        return xml
    chunks = []
    for rownum, cells, style in rows:
        if not cells:
            continue
        cols = sorted(cells)
        spans = f'{cols[0]}:{cols[-1]}'
        parts = [f'<row r="{rownum}" spans="{spans}">']
        for col in cols:
            text = cells[col]
            if text is None or text == "":
                continue
            ref = f"{get_column_letter(col)}{rownum}"
            s = f' s="{style}"' if style is not None else ""
            parts.append(
                f'<c r="{ref}"{s} t="inlineStr">'
                f'<is><t xml:space="preserve">{xml_escape(str(text))}</t></is></c>'
            )
        parts.append("</row>")
        chunks.append("".join(parts))
    blob = "".join(chunks)
    if "</sheetData>" not in xml:
        raise ValueError("no </sheetData> in sheet XML")
    return xml.replace("</sheetData>", blob + "</sheetData>", 1)


def bump_dimension(xml: str, last_row: int) -> str:
    m = re.search(r'<dimension ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"/>', xml)
    if not m:
        return xml
    c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
    if last_row <= r2:
        return xml
    return xml[: m.start()] + f'<dimension ref="{c1}{r1}:{c2}{last_row}"/>' + xml[m.end():]


def real_method(signature: str) -> str | None:
    m = METHOD_RE.search((signature or "").strip())
    return m.group(1) if m else None


def utcid_columns(ws) -> dict[str, int]:
    out = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(UTCID_HEADER_ROW, c).value
        if isinstance(v, str) and v.strip().startswith("UTCID"):
            out[v.strip()] = c
    return out


def label_row_style(ws) -> tuple[int, str | None]:
    """Last labelled row in column B, and that label cell's style index (for visual match)."""
    last, style = UTCID_HEADER_ROW, None
    for r in range(1, ws.max_row + 1):
        cell = ws.cell(r, LABEL_COL)
        if isinstance(cell.value, str) and cell.value.strip():
            last = r
            style = cell.style_id if hasattr(cell, "style_id") else None
    return last, style


def manual_marks(ws, cols: dict[str, int]) -> Counter:
    """Count the sheet's OWN manual Passed/Failed marks. Read-only."""
    row = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, LABEL_COL).value
        if isinstance(v, str) and "Passed/Failed" in v:
            row = r
            break
    tally = Counter()
    if row is None:
        return tally
    for c in cols.values():
        v = ws.cell(row, c).value
        v = str(v).strip().upper() if v is not None else ""
        tally["P" if v.startswith("P") else "F" if v.startswith("F") else "U"] += 1
    return tally


def sheet_xml_map(zf: zipfile.ZipFile) -> dict[str, str]:
    """sheet display name -> zip entry path for its worksheet XML."""
    wbxml = zf.read("xl/workbook.xml").decode("utf-8")
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rmap = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    out = {}
    for name, rid in re.findall(r'<sheet name="([^"]+)"[^>]*r:id="(rId\d+)"', wbxml):
        tgt = rmap.get(rid, "")
        if tgt:
            out[name.replace("&amp;", "&")] = "xl/" + tgt.lstrip("/")
    return out


def main() -> int:
    matrix = json.loads(MATRIX.read_text(encoding="utf-8"))
    results = json.loads(RESULTS.read_text(encoding="utf-8")) if RESULTS.exists() else []
    symbols = json.loads(SYMBOLS.read_text(encoding="utf-8"))

    by_sheet_bucket: dict[str, Counter] = defaultdict(Counter)
    for r in matrix:
        by_sheet_bucket[r["sheet"]][r["bucket"]] += 1
    res_index = {(r["sheet"], r["utcid"]): r for r in results}

    # ---- READ-ONLY pass over the ORIGINAL to decide every edit --------------------------
    wb = openpyxl.load_workbook(SRC)          # never saved; SRC is untouched
    with zipfile.ZipFile(SRC) as zf:
        xmlmap = sheet_xml_map(zf)

    edits: dict[str, dict] = defaultdict(lambda: {"num": {}, "str": {}, "rows": [], "last": 0})
    changelog: list[str] = []
    filled = 0

    # 1. additive per-UTCID rows on every function sheet
    for name in wb.sheetnames:
        if name in META_SHEETS:
            continue
        ws = wb[name]
        cols = utcid_columns(ws)
        if not cols:
            continue
        base, style = label_row_style(ws)
        r_auto, r_align = base + 1, base + 2
        auto_cells: dict[int, str] = {LABEL_COL: ROW_AUTOMATED}
        algn_cells: dict[int, str] = {LABEL_COL: ROW_ALIGNMENT}
        for utcid, col in cols.items():
            hit = res_index.get((name, utcid))
            if not hit:
                continue                     # BLANK -- never "Failed", never "Untested"
            auto_cells[col] = hit["test_name"]
            algn_cells[col] = hit.get("spec_alignment") or ""
            filled += 2
        path = xmlmap[name]
        edits[path]["rows"] = [(r_auto, auto_cells, style), (r_align, algn_cells, style)]
        edits[path]["last"] = r_align

    # 2. Statistics tally correction (+ refreshed cached sums on the =SUM() row)
    fn = wb["Functions"]
    code2sheet = {}
    for r in range(FUNC_FIRST, FUNC_LAST + 1):
        code, sheet = fn.cell(r, F_CODE).value, fn.cell(r, F_SHEET).value
        if code and sheet:
            code2sheet[str(code).strip()] = str(sheet).strip()

    st = wb["Statistics"]
    stats_path = xmlmap["Statistics"]
    totals = Counter()
    for r in range(STATS_FIRST, STATS_LAST + 1):
        code = st.cell(r, 2).value
        if code is None:
            continue
        sheet = code2sheet.get(str(code).strip())
        if sheet is None or sheet not in by_sheet_bucket:
            changelog.append(f"Statistics!B{r}: function code {code!r} unresolved -- LEFT UNCHANGED")
            continue
        b = by_sheet_bucket[sheet]
        marks = manual_marks(wb[sheet], utcid_columns(wb[sheet]))
        want = {
            C_PASSED: marks["P"], C_FAILED: marks["F"], C_UNTESTED: marks["U"],
            C_N: b["N"], C_A: b["A"], C_B: b["B"], C_TOTAL: b["N"] + b["A"] + b["B"],
        }
        for col, new in want.items():
            totals[col] += new
            old = st.cell(r, col).value
            if old != new:
                ref = f"{get_column_letter(col)}{r}"
                edits[stats_path]["num"][ref] = new
                changelog.append(f"Statistics!{ref} ({sheet} / {code}): {old} -> {new}")
    # refresh the cached values behind the (preserved) =SUM() formulas
    for col in (C_PASSED, C_FAILED, C_UNTESTED, C_N, C_A, C_B, C_TOTAL):
        ref = f"{get_column_letter(col)}{STATS_SUM}"
        old = st.cell(STATS_SUM, col).value
        if isinstance(old, str) and old.startswith("="):
            old = None
        edits[stats_path]["num"][ref] = totals[col]
        changelog.append(f"Statistics!{ref} (=SUM cached value): {old} -> {totals[col]}")

    # 3. Functions sheet: the 11 MISMAPPED symbols
    func_path = xmlmap["Functions"]
    mismapped = {s["sheet"]: s for s in symbols if s["status"] == "MISMAPPED"}
    for r in range(FUNC_FIRST, FUNC_LAST + 1):
        sheet = fn.cell(r, F_SHEET).value
        if not sheet:
            continue
        sym = mismapped.get(str(sheet).strip())
        if not sym:
            continue
        method = real_method(sym.get("signature", ""))
        if not method:
            changelog.append(f"Functions!row {r} ({sheet}): signature unparseable -- LEFT UNCHANGED")
            continue
        for col, new in ((F_CLASS, sym["actualClass"]), (F_METHOD, f"{method}()")):
            old = fn.cell(r, col).value
            if old != new:
                ref = f"{get_column_letter(col)}{r}"
                edits[func_path]["str"][ref] = new
                changelog.append(f"Functions!{ref} ({sheet}): {old!r} -> {new!r}")
    wb.close()

    # ---- APPLY: copy the zip entry-by-entry, rewriting only what must change ------------
    modified = set(edits) | {"xl/workbook.xml"}
    DST.parent.mkdir(parents=True, exist_ok=True)
    tmp = DST.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(SRC) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in edits:
                xml = data.decode("utf-8")
                spec = edits[item.filename]
                for ref, val in spec["num"].items():
                    xml = set_numeric(xml, ref, val)
                for ref, text in spec["str"].items():
                    xml = set_inline_string(xml, ref, text)
                if spec["rows"]:
                    xml = append_rows(xml, spec["rows"])
                    xml = bump_dimension(xml, spec["last"])
                data = xml.encode("utf-8")
            elif item.filename == "xl/workbook.xml":
                xml = data.decode("utf-8")
                if "fullCalcOnLoad" not in xml:
                    xml = re.sub(r"<calcPr([^>]*?)/>", r'<calcPr\1 fullCalcOnLoad="1"/>', xml, count=1)
                data = xml.encode("utf-8")
            zi = zipfile.ZipInfo(item.filename, date_time=item.date_time)
            zi.compress_type = item.compress_type
            zi.external_attr = item.external_attr
            zi.internal_attr = item.internal_attr
            zi.create_system = item.create_system
            zout.writestr(zi, data)
    tmp.replace(DST)

    # ---- VERIFICATION GATE -------------------------------------------------------------
    with zipfile.ZipFile(SRC) as a, zipfile.ZipFile(DST) as b:
        na, nb = set(a.namelist()), set(b.namelist())
        lost, added = sorted(na - nb), sorted(nb - na)
        drift = [n for n in sorted(na & nb) if n not in modified and a.read(n) != b.read(n)]
        print("ZIP INTEGRITY")
        print(f"  entries: original {len(na)}  ->  v2 {len(nb)}")
        print(f"  lost entries            : {lost if lost else 'NONE'}")
        print(f"  unexpected new entries  : {added if added else 'NONE'}")
        print(f"  modified (intended)     : {len(modified & nb)} -> {sorted(modified & nb)}")
        print(f"  byte-drift in untouched : {drift if drift else 'NONE'}")
        ok = not lost and not added and not drift
        print(f"  VERDICT: {'PASS' if ok else 'FAIL'}")
    if lost or added or drift:
        return 1

    print()
    print(f"wrote {DST.relative_to(REPO)}")
    print(f"  original untouched: {SRC.relative_to(REPO)}")
    print(f"  automated-result cells filled : {filled}  ({len(results)} UTCIDs x 2)")
    print(f"  corrections made              : {len(changelog)}")
    for line in changelog[:6]:
        print(f"    {line}")
    if len(changelog) > 6:
        print(f"    ... {len(changelog) - 6} more (see docs/audit/xlsx-changelog.md)")

    CHANGELOG_RAW.write_text("\n".join(changelog) + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
