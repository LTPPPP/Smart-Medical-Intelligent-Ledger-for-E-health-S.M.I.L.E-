#!/usr/bin/env python3
"""Parse Report5_Unit Test.xlsx into docs/audit/uc-matrix.json.

One object per UTCID. Target symbols are resolved from docs/audit/qa-recon-symbols.json,
NOT from the spreadsheet's own Class/Function columns -- for the 11 MISMAPPED rows the
spreadsheet names the wrong class, so claimedClass/claimedMethod must be ignored and the
real target read from actualClass + signature + file:line.

Read-only with respect to the workbook. Deterministic: same input -> same output.

Usage: python3 scripts/qa/parse-uc-matrix.py
"""

import json
import re
import sys
from collections import Counter
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
XLSX = REPO / "docs/OneDrive_1_8-1-2026/Report5_Unit Test.xlsx"
SYMBOLS = REPO / "docs/audit/qa-recon-symbols.json"
OUT = REPO / "docs/audit/uc-matrix.json"

META_SHEETS = {"Cover", "Functions", "Statistics"}

# Expected totals, independently recounted twice. The Statistics sheet claims 431 (A=332);
# that 33-case gap is a roll-up arithmetic error in the workbook. 398 is the truth.
EXPECT_TOTAL = 398
EXPECT_BUCKETS = {"N": 87, "A": 299, "B": 12}

SERVICE_BY_DIR = {
    "iam-service": "iam",
    "clinical-emr-service": "clinical-emr",
    "payment-service": "payment",
    "gateway-service": "gateway",
}

UTCID_RE = re.compile(r"UTCID\d+")
# Pulls the real method name out of a signature like "async initiate( dto: ... ): Promise<"
METHOD_RE = re.compile(r"(?:async\s+)?([A-Za-z_][A-Za-z0-9_]*)\s*\(")


def norm(v):
    """Normalise a cell value to a stripped string, or None if blank."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def service_for(file_path):
    for d, name in SERVICE_BY_DIR.items():
        if f"/{d}/" in file_path:
            return name
    return "unknown"


def load_symbols():
    """sheet name -> resolved real target symbol."""
    rows = json.loads(SYMBOLS.read_text())
    by_sheet = {}
    for r in rows:
        sig = r.get("signature") or ""
        m = METHOD_RE.search(sig)
        if not m:
            # Fall back to the claimed method only for EXACT rows, where it is trustworthy.
            if r.get("status") == "EXACT":
                method = (r.get("claimedMethod") or "").rstrip("()")
            else:
                raise SystemExit(
                    f"cannot derive method name for MISMAPPED sheet {r['sheet']!r}"
                )
        else:
            method = m.group(1)
        by_sheet[r["sheet"]] = {
            "targetClass": r.get("actualClass"),
            "targetMethod": method,
            "file": r.get("file"),
            "line": r.get("line"),
            "service": service_for(r.get("file") or ""),
            "symbolStatus": r.get("status"),
        }
    return by_sheet


def find_header(ws):
    """Locate the UTCID header row -> (row_index, {utcid: column_index})."""
    for r in range(1, ws.max_row + 1):
        cols = {}
        for c in range(1, ws.max_column + 1):
            v = norm(ws.cell(r, c).value)
            if v and UTCID_RE.fullmatch(v):
                cols[v] = c
        if cols:
            return r, cols
    return None, None


def parse_sheet(ws):
    """-> (list of utcids in order, {utcid: {...}})."""
    hr, cols = find_header(ws)
    if not hr:
        raise ValueError("no UTCID header row")

    utcids = list(cols.keys())
    data = {
        u: {"inputs": {}, "expected_detail": {}, "bucket": None,
            "sheet_passed_failed": None, "sheet_executed_date": None}
        for u in utcids
    }

    section = None
    group = None

    for r in range(hr + 1, ws.max_row + 1):
        a = norm(ws.cell(r, 1).value)
        b = norm(ws.cell(r, 2).value)
        c = norm(ws.cell(r, 3).value)

        if a:
            section = a  # 'Condition' | 'Confirm' | 'Result'
        if b:
            group = b

        # --- Result section: per-UTCID scalar rows, keyed off column B ---
        if section == "Result" and b:
            if b.startswith("Type("):
                for u, col in cols.items():
                    data[u]["bucket"] = norm(ws.cell(r, col).value)
                continue
            if b.startswith("Passed/Failed"):
                for u, col in cols.items():
                    data[u]["sheet_passed_failed"] = norm(ws.cell(r, col).value)
                continue
            if b.startswith("Executed Date"):
                for u, col in cols.items():
                    data[u]["sheet_executed_date"] = norm(ws.cell(r, col).value)
                continue
            continue

        # --- Condition / Confirm: column C holds a candidate value, 'O' marks which
        #     UTCIDs take it. A row may both open a group (col B) and carry a value (col C).
        if c is None or group is None:
            continue

        for u, col in cols.items():
            mark = norm(ws.cell(r, col).value)
            if not mark:
                continue
            bucket = "inputs" if section == "Condition" else "expected_detail"
            slot = data[u][bucket]
            # A group can legitimately be marked on several rows; keep every value.
            if group in slot:
                prev = slot[group]
                slot[group] = prev + [c] if isinstance(prev, list) else [prev, c]
            else:
                slot[group] = c

    return utcids, data


def expected_summary(detail):
    """Flatten the Confirm section into one human-readable string."""
    parts = []
    for key in ("Return", "Exception", "Log message"):
        if key in detail:
            v = detail[key]
            v = " ; ".join(v) if isinstance(v, list) else v
            parts.append(f"{key}: {v}")
    for key, v in detail.items():
        if key not in ("Return", "Exception", "Log message"):
            v = " ; ".join(v) if isinstance(v, list) else v
            parts.append(f"{key}: {v}")
    return " | ".join(parts)


def main():
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl required: pip install --break-system-packages openpyxl")

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    symbols = load_symbols()

    sheets = [n for n in wb.sheetnames if n not in META_SHEETS]
    if len(sheets) != 87:
        sys.exit(f"expected 87 function sheets, found {len(sheets)}")

    matrix = []
    per_sheet = {}
    missing_symbol = []

    for name in sheets:
        ws = wb[name]
        utcids, data = parse_sheet(ws)
        per_sheet[name] = len(utcids)

        sym = symbols.get(name)
        if sym is None:
            missing_symbol.append(name)
            sym = {"targetClass": None, "targetMethod": None, "file": None,
                   "line": None, "service": "unknown", "symbolStatus": "NOT_IN_SYMBOL_MAP"}

        # Function Name cell (col K, row 2) = the spreadsheet's own claim; kept for audit.
        claimed = norm(ws.cell(2, 11).value)

        for u in utcids:
            d = data[u]
            matrix.append({
                "sheet": name,
                "utcid": u,
                "bucket": d["bucket"],
                "inputs": d["inputs"],
                "expected": expected_summary(d["expected_detail"]),
                "expected_detail": d["expected_detail"],
                "targetClass": sym["targetClass"],
                "targetMethod": sym["targetMethod"],
                "file": sym["file"],
                "line": sym["line"],
                "service": sym["service"],
                "symbolStatus": sym["symbolStatus"],
                "spreadsheetClaimedSymbol": claimed,
                "sheetReportedResult": d["sheet_passed_failed"],
                "sheetReportedDate": d["sheet_executed_date"],
            })

    total = len(matrix)
    buckets = Counter(m["bucket"] for m in matrix)

    print(f"function sheets      : {len(sheets)}")
    print(f"total UTCIDs         : {total}")
    print(f"buckets              : N={buckets.get('N')} A={buckets.get('A')} B={buckets.get('B')}")
    if missing_symbol:
        print(f"sheets w/o symbol map: {missing_symbol}")

    ok = total == EXPECT_TOTAL and all(buckets.get(k) == v for k, v in EXPECT_BUCKETS.items())
    if not ok:
        print("\nASSERTION FAILED -- expected total=398, N=87, A=299, B=12")
        print("per-sheet counts:")
        for n, cnt in sorted(per_sheet.items(), key=lambda kv: -kv[1]):
            print(f"  {cnt:3d}  {n}")
        sys.exit(1)

    print("ASSERTION PASSED     : total=398, N=87, A=299, B=12")

    by_service = Counter(m["service"] for m in matrix)
    print("\nper-service UTCID breakdown:")
    for s in ("iam", "clinical-emr", "payment", "gateway", "unknown"):
        if by_service.get(s):
            sheets_in = len({m["sheet"] for m in matrix if m["service"] == s})
            print(f"  {s:14} {by_service[s]:4d} UTCIDs across {sheets_in:3d} sheets")

    OUT.write_text(json.dumps(matrix, indent=2, ensure_ascii=False) + "\n")
    print(f"\nwrote {OUT.relative_to(REPO)} ({total} rows)")


if __name__ == "__main__":
    main()
